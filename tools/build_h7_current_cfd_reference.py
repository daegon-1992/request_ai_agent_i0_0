#!/usr/bin/env python3
"""Build a code-derived reference pack for the current h7_v0 CFD workflow."""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from build_structural_sme_request_pack import (  # noqa: E402
    NS_MAIN,
    REVIEW_STATUS,
    SheetSpec,
    Validation,
    _add_bullets,
    _add_table,
    write_xlsx,
)
from request_ai_agent_i0_0.analysis_type_master import ANALYSIS_TYPE_MASTER  # noqa: E402
from request_ai_agent_i0_0.condition_fieldsets import (  # noqa: E402
    build_condition_fieldset,
    get_active_case_matrix_columns,
)
from request_ai_agent_i0_0.constants import (  # noqa: E402
    ANALYSIS_TYPE_OPTIONS,
    DISABLED_ANALYSIS_TYPE_OPTIONS,
    ENABLED_ANALYSIS_TYPE_OPTIONS,
)
from request_ai_agent_i0_0.schema import (  # noqa: E402
    ANALYSIS_OVERVIEW_SPECS,
    BASIC_INFO_SPECS,
)


OUTPUT_DIR = ROOT / "docs" / "structural_analysis_sme_request"
MARKDOWN_PATH = OUTPUT_DIR / "04_h7_v0_현행_유동해석_기준.md"
DOCX_PATH = OUTPUT_DIR / "h7_v0_현행_유동해석_기준서.docx"
XLSX_PATH = OUTPUT_DIR / "h7_v0_현행_유동해석_Field_Master.xlsx"
REFERENCE_VERSION = "h7_v0-current-baseline-20260820"

EXPECTED_ACTIVE_TYPES = (
    "풍량",
    "기류 패턴",
    "이슬맺힘",
    "열교환기 유속 프로파일",
)
EXPECTED_FIELD_KEYS = {
    "풍량": ("fan_rpm", "name", "tube_diameter", "fin_type", "row_count", "fpi"),
    "기류 패턴": (
        "fan_rpm", "name", "tube_diameter", "fin_type", "row_count", "fpi",
        "heat_exchanger_temp", "room_temp",
    ),
    "이슬맺힘": (
        "fan_rpm", "name", "tube_diameter", "fin_type", "row_count", "fpi",
        "heat_exchanger_temp", "heat_exchanger_rh", "room_temp", "room_rh",
    ),
    "열교환기 유속 프로파일": ("fan_rpm", "name", "tube_diameter", "fin_type", "row_count", "fpi"),
}
EXPECTED_SHEETS = (
    "작성안내",
    "01_해석유형_우선순위",
    "02_완료사례_목록",
    "03_Field_Rule_Master",
    "04_핵심객체_규칙",
    "05_Case_조합규칙",
    "06_용어_안내_예외",
    "07_검수_승인",
    "08_변경이력",
    "09_워크숍_기록",
    "10_현행_Gap_목록",
)


RESULT_LABELS = {
    "풍량": "풍량, 유동 분포",
    "기류 패턴": "기류 경로, 공간 분포",
    "이슬맺힘": "결로 위험 위치, 온·습도 조건 비교",
    "열교환기 유속 프로파일": "열교환기 면 유속 분포, 균일도",
    "기류도달거리": "도달 거리, 기류 분포",
    "PDB": "발열, 냉각 성능",
    "실사용 해석": "실사용 유동·열 성능",
    "집진해석(먼지거동)": "입자 이동, 포집 거동",
    "PCB발열": "PCB 온도 분포, 발열 영향",
    "다상유동": "계면, 혼합·이동 거동",
}

FIELD_HELP = {
    "request_no": ("시스템이 발급하는 의뢰 식별번호", "시스템에서 자동 생성합니다."),
    "division": ("의뢰자가 소속된 사업부", "소속 사업부를 선택해 주세요."),
    "department": ("의뢰자 소속 부서", "소속 부서를 입력해 주세요."),
    "requester_name": ("해석 의뢰 담당자 이름", "요청자 이름을 입력해 주세요."),
    "requester_role": ("의뢰자의 직급", "직급을 선택하거나 입력해 주세요."),
    "project_name": ("해석을 수행할 PMS 프로젝트", "프로젝트명(PMS)을 선택하거나 입력해 주세요."),
    "development_grade": ("제품 개발 등급", "개발 등급을 선택해 주세요."),
    "npi_stage": ("제품 개발 진행 단계", "NPI 단계를 선택해 주세요."),
    "model_suffix": ("해석 대상 모델명 또는 Suffix", "모델명(Model Suffix)을 입력해 주세요."),
    "request_date": ("해석 의뢰 요청일", "의뢰 요청일을 확인해 주세요."),
    "desired_completion_date": ("결과를 받고 싶은 날짜", "희망 완료일을 선택해 주세요."),
    "request_description": ("해석을 요청하게 된 문제와 배경", "해석을 요청하게 된 배경을 설명해 주세요."),
    "decision_use": ("해석 결과를 사용할 의사결정", "결과를 어디에 활용할지 선택해 주세요."),
    "additional_result_request": ("해석으로 확인하고 싶은 구체적인 내용", "해석으로 확인하고 싶은 내용을 입력해 주세요."),
}


def _active_groups(analysis_type: str) -> str:
    fieldset = build_condition_fieldset({"analysis_type": analysis_type})
    return ", ".join(group["label"] for group in fieldset["groups"] if group["active"])


def _active_fields(analysis_type: str) -> tuple[str, ...]:
    return tuple(build_condition_fieldset({"analysis_type": analysis_type})["field_keys"])


def _analysis_rows() -> tuple[tuple[str, ...], ...]:
    rows = []
    for analysis_type in ANALYSIS_TYPE_OPTIONS:
        enabled = analysis_type in ENABLED_ANALYSIS_TYPE_OPTIONS
        rows.append((
            analysis_type,
            ANALYSIS_TYPE_MASTER.get(analysis_type, "현행 안내 없음"),
            "기존 사업부·제품군·Platform 범위",
            RESULT_LABELS.get(analysis_type, "담당자 확인 필요"),
            "현행 코드에 부적합 조건 정의 없음",
            "현행 자료 없음",
            "구현됨" if enabled else "미정",
            "현행 자료 없음",
            "현행 지원" if enabled else "향후(예정)",
            "선택 가능" if enabled else "예정·선택 불가",
            _active_groups(analysis_type) if enabled else "유형별 fieldset 미정",
            "코드 기준 정리 완료" if enabled else "검토 필요",
            "constants.py; condition_fieldsets.py; analysis_type_master.py",
        ))
    return tuple(rows)


def _case_reference_rows() -> tuple[tuple[str, ...], ...]:
    return (
        (
            "REF-CFD-AIRVOLUME-01", "풍량", "코드 동작 예시(실제 완료사례 아님)", "단일 하중·단일 형상",
            "Base 제품의 풍량과 유동 분포 확인", "Base 총조립 제품", "Base 도면번호 1건",
            "현행 입력 없음", "운전 1의 팬 RPM", "현행 입력 없음", "열교환기 사양 1",
            "현행 실제 재문의 자료 없음", "형상 1+운전 1+사양 1", "Case 1",
            "풍량과 유동 분포", "정량 합격 기준은 현행 코드에 없음", "현행 실제 보고서 없음",
            "해당 없음", "검토 필요", "업무 사례 3건으로 교체 필요",
        ),
        (
            "REF-CFD-AIRFLOW-01", "기류 패턴", "코드 동작 예시(실제 완료사례 아님)", "입력 누락·재문의",
            "제품 주변 기류 경로 확인", "Base 총조립 제품", "Base 도면번호 1건",
            "현행 입력 없음", "운전 1의 팬 RPM", "현행 입력 없음", "열교환기 사양 1",
            "온도 조건이 없으면 `없음` 선택 안내", "공간 온도 없음, 취출 온도 없음",
            "형상 1+운전 1+사양 1+공간 온도 없음+취출 온도 없음", "기류 경로와 공간 분포",
            "정량 합격 기준은 현행 코드에 없음", "현행 실제 보고서 없음", "해당 없음", "검토 필요",
            "`없음`은 코드가 허용하는 입력이며 업무 적합성 재확인 필요",
        ),
        (
            "REF-CFD-DEW-01", "이슬맺힘", "코드 동작 예시(실제 완료사례 아님)", "다중 Case",
            "온·습도 조건별 결로 위험 비교", "Base 및 비교 총조립 제품", "서로 다른 도면번호 2건과 차이 설명",
            "현행 입력 없음", "운전 조건 1개 이상", "현행 입력 없음", "열교환기 사양 1개 이상",
            "현행 실제 재문의 자료 없음", "공간·취출 온도와 상대습도 입력",
            "형상별로 온·습도 조건을 수동 매핑", "결로 위험 위치와 조건 비교",
            "정량 판정 기준은 현행 코드에 없음", "현행 실제 보고서 없음", "해당 없음", "검토 필요",
            "업무 사례와 결로 판정 기준 필요",
        ),
        (
            "REF-CFD-HEX-01", "열교환기 유속 프로파일", "코드 동작 예시(실제 완료사례 아님)", "다중 Case",
            "열교환기 사양별 면 유속 균일도 비교", "Base 총조립 제품", "Base 도면번호 1건",
            "현행 입력 없음", "운전 1과 운전 2", "현행 입력 없음", "사양 1과 사양 2",
            "현행 실제 재문의 자료 없음", "각 운전·사양 조합을 Case로 선택",
            "필요한 조합만 수동 추가", "열교환기 면 유속 분포와 균일도",
            "균일도 허용 기준은 현행 코드에 없음", "현행 실제 보고서 없음", "해당 없음", "검토 필요",
            "모든 입력 조건은 최소 한 Case에서 사용되어야 Word 생성 가능",
        ),
        (
            "GAP-CFD-ACTUAL", "전체 활성 유형", "자료 Gap", "실제 완료사례",
            "저장소에 최초 사용자 요청 자료 없음", "현행 자료 없음", "현행 자료 없음", "현행 자료 없음",
            "현행 자료 없음", "현행 자료 없음", "현행 자료 없음", "현행 자료 없음", "현행 자료 없음",
            "현행 자료 없음", "현행 자료 없음", "현행 자료 없음", "현행 자료 없음", "검토 필요",
            "검토 필요", "유동해석 담당자로부터 유형별 최소 3건 수집 필요",
        ),
    )


def _common_field_rows() -> list[tuple[str, ...]]:
    rows: list[tuple[str, ...]] = []

    def add(
        rule_id: str,
        group: str,
        label: str,
        definition: str,
        analysis_types: str,
        required: str,
        activation: str,
        owner: str,
        input_type: str,
        unit: str,
        choices: str,
        repeatable: str,
        binding: str,
        unknown: str,
        blocking: str,
        question: str,
        normal: str,
        error: str,
        status: str,
        source: str,
    ) -> None:
        rows.append((
            rule_id, group, label, definition, analysis_types, required, activation, owner, input_type,
            unit, choices, repeatable, binding, unknown, blocking, question, normal, error, status, source,
        ))

    context_rows = (
        ("business_unit", "사업부", "의뢰 대상 제품의 사업부", "필수", "선택", "제품 분류 Master", "사업부를 선택해 주세요."),
        ("product_group", "제품군", "해석 대상 제품군", "필수", "선택", "선택 사업부", "제품군을 선택해 주세요."),
        ("platform", "Platform", "해석 대상 제품 Platform", "필수", "선택", "선택 제품군", "Platform을 선택해 주세요."),
        ("analysis_type", "해석유형", "요청 목적에 맞는 유동해석 유형", "필수", "선택", "활성 4종", "해석유형을 선택해 주세요."),
        ("operation_mode", "운전 모드", "향후 운전 분류용 Context 필드", "현행 미사용", "텍스트", "현행 선택지 없음", "현행 별도 질문 없음"),
        ("context_locked", "Context 확정 여부", "제품 분류와 해석유형의 확정 상태", "시스템 필수", "예/아니오", "시스템", "선택 내용을 확정할까요?"),
    )
    for key, label, definition, required, input_type, choices, question in context_rows:
        add(
            f"request_context.{key}", "공통", label, definition, "전체", required, "Context 확정 단계",
            "시스템" if key == "context_locked" else "의뢰자", input_type, "", choices, "아니오",
            "request_context", "확정 전 모름은 미입력", "필수 4개가 없으면 Context 확정 차단" if required == "필수" else "현행 차단 없음",
            question, "값 입력 또는 선택", "빈 값", "코드 기준 정리 완료", "constants.py; request_context_confirmation.py",
        )

    for spec in (*BASIC_INFO_SPECS, *ANALYSIS_OVERVIEW_SPECS):
        definition, question = FIELD_HELP.get(spec.key, (spec.label, f"{spec.label}을(를) 입력해 주세요."))
        required = "필수" if spec.required else ("시스템 생성" if spec.key == "request_no" else "선택")
        choices = {
            "development_grade": "A, B, Ca, Cb, Cc, 선행, 미정, 직접 입력",
            "npi_stage": "CP, DV, PV, MP, 미정, 직접 입력",
            "project_name": "미정, 직접 입력",
            "model_suffix": "미정, 직접 입력",
        }.get(spec.key, "직접 입력")
        input_type = "날짜" if spec.value_type == "date" else ("선택+직접 입력" if spec.key in {"development_grade", "npi_stage", "project_name", "model_suffix"} else "텍스트")
        add(
            f"{spec.section}.{spec.key}", "의뢰 정보" if spec.section == "basic_info" else "해석 개요",
            spec.label, definition, "전체", required, "항상", "시스템" if spec.key in {"request_no", "request_date"} else "의뢰자",
            input_type, "", choices, "아니오", spec.section, "미정 선택은 일부 필드에서 허용",
            "누락 시 제출 차단" if spec.required else "현행 차단 없음", question, "유효한 값", "빈 값",
            "코드 기준 정리 완료", "constants.py; schema.py; validator.py",
        )

    geometry_rows = (
        ("geometry.base_product", "Base 제품", "기준이 되는 총조립 제품", "필수", "객체", "1개 고정", "아니오", "형상 축", "Base 제품 부재 시 차단", "기준 제품을 등록해 주세요."),
        ("geometry.drawing_no", "도면번호 (NPDM MCAD)", "해석할 총조립 형상의 식별번호", "필수", "텍스트", "영문·숫자·하이픈", "제품별 반복", "제품", "누락·형식 오류·중복 시 차단", "총조립 형상의 도면번호를 입력해 주세요."),
        ("geometry.comparison_products", "비교 제품", "Base와 비교할 추가 총조립 제품", "선택", "반복 행", "0개 이상", "예", "형상 축", "추가한 제품의 도면번호 누락 시 차단", "비교할 제품이 있으면 추가해 주세요."),
        ("geometry.difference_from_base", "Base 대비 차이", "비교 제품에 반영된 형상·조립 상태 차이", "조건부 필수", "텍스트", "직접 입력", "제품별 반복", "비교 제품", "비교 제품에서 누락 시 차단", "Base 제품과 무엇이 다른지 설명해 주세요."),
    )
    for key, label, definition, required, input_type, choices, repeatable, binding, blocking, question in geometry_rows:
        add(
            key, "제품/형상", label, definition, "전체", required, "항상" if required != "조건부 필수" else "비교 제품을 추가한 경우",
            "의뢰자", input_type, "", choices, repeatable, binding, "임시 도면번호 허용 안내", blocking,
            question, "서로 다른 유효 도면번호", "중복 또는 설명 누락", "코드 기준 정리 완료", "state.py; validator.py; ui.py",
        )

    all_types = ", ".join(EXPECTED_ACTIVE_TYPES)
    condition_rows = (
        ("conditions.operating.name", "운전명", "반복 운전 조건의 시스템 이름", all_types, "시스템 필수", "운전 카드", "시스템", "텍스트", "", "운전 N", "예", "Case 운전 열", "현행 해당 없음", "시스템 생성", "운전 조건 이름", "운전 1", "수정 대상 아님"),
        ("conditions.operating.fan_count", "팬 개수", "운전 조건에 포함되는 팬 수", all_types, "구성 필드", "운전 카드", "의뢰자", "선택+직접 입력", "개", "1~4 또는 직접 입력", "예", "팬 구성", "현행 미정", "핵심 validator 차단 없음", "팬은 몇 개인가요?", "2개", "0 또는 비정상 입력"),
        ("conditions.operating.fan_location", "팬 위치", "여러 팬을 구분하는 위치명", all_types, "조건부 안내", "팬 2개 이상", "의뢰자", "텍스트", "", "직접 입력", "예", "팬 구성", "현행 미정", "핵심 validator 차단 없음", "각 팬의 위치를 입력해 주세요.", "상부/하부", "여러 팬의 위치 미입력"),
        ("conditions.operating.fan_rpm", "팬 회전수(RPM)", "각 팬의 운전 회전수", all_types, "필수", "운전 카드", "의뢰자", "숫자/텍스트", "RPM", "물리 범위 미정", "예", "Case 운전 열", "현행 통합 정책 없음", "활성 팬 RPM 누락 시 차단", "각 팬의 회전수를 입력해 주세요.", "900 RPM", "빈 RPM"),
        ("conditions.heat_exchanger.name", "사양명", "반복 열교환기 사양의 시스템 이름", all_types, "시스템 필수", "열교환기 카드", "시스템", "텍스트", "", "사양 N", "예", "Case 사양 열", "해당 없음", "시스템 생성", "열교환기 사양 이름", "사양 1", "수정 대상 아님"),
        ("conditions.heat_exchanger.type", "열교환기 유형", "Fin&Tube 또는 Micro-Channel 구분", all_types, "구성 필드", "열교환기 카드", "의뢰자/시스템", "선택", "", "Fin&Tube, Micro-Channel", "예", "열교환기 사양", "현행 미정", "핵심 validator 차단 없음", "열교환기 유형을 선택해 주세요.", "Fin&Tube", "지원하지 않는 유형"),
        ("conditions.heat_exchanger.tube_diameter", "관 직경(Pi)", "열교환기 관 직경 또는 Catalog 구분값", all_types, "필수", "열교환기 카드", "의뢰자", "텍스트", "mm", "W 접두값이면 Micro-Channel", "예", "열교환기 사양", "현행 통합 정책 없음", "누락 시 차단", "관 직경(Pi)을 입력해 주세요.", "7.0 mm", "빈 값"),
        ("conditions.heat_exchanger.fin_type", "Fin type", "열교환기 Fin 형상 유형", all_types, "필수", "열교환기 카드", "의뢰자/시스템", "텍스트", "", "Micro-Channel이면 Flat", "예", "열교환기 사양", "현행 통합 정책 없음", "누락 시 차단", "Fin type을 입력해 주세요.", "Slit", "빈 값"),
        ("conditions.heat_exchanger.row_count", "열 수", "열교환기 관 배열의 열 수", all_types, "필수", "열교환기 카드", "의뢰자", "숫자/텍스트", "", "물리 범위 미정", "예", "열교환기 사양", "현행 통합 정책 없음", "누락 시 차단", "열 수를 입력해 주세요.", "2", "빈 값"),
        ("conditions.heat_exchanger.fpi", "FPI", "인치당 Fin 수", all_types, "필수", "열교환기 카드", "의뢰자", "숫자/텍스트", "FPI", "물리 범위 미정", "예", "열교환기 사양", "현행 통합 정책 없음", "누락 시 차단", "FPI를 입력해 주세요.", "18", "빈 값"),
        ("conditions.supply_air.temp", "취출 온도", "열교환기 이후 취출 공기 온도", "기류 패턴, 이슬맺힘", "필수", "취출 공기 카드", "의뢰자", "숫자/선택", "°C", "기류 패턴은 없음 선택 가능", "예", "Case 취출 온도 열", "기류 패턴에서 없음 허용", "누락 시 차단", "취출 온도를 입력해 주세요.", "15 °C", "빈 값"),
        ("conditions.supply_air.rh", "취출 상대습도", "취출 공기의 상대습도", "이슬맺힘", "필수", "이슬맺힘", "의뢰자", "숫자", "%", "물리 범위 미정", "예", "Case 취출 RH 열", "현행 통합 정책 없음", "누락 시 차단", "취출 상대습도를 입력해 주세요.", "80%", "빈 값"),
        ("conditions.space.temp", "공간 온도", "제품 설치 공간의 공기 온도", "기류 패턴, 이슬맺힘", "필수", "공간 환경 카드", "의뢰자", "숫자/선택", "°C", "기류 패턴은 없음 선택 가능", "예", "Case 공간 온도 열", "기류 패턴에서 없음 허용", "누락 시 차단", "공간 온도를 입력해 주세요.", "27 °C", "빈 값"),
        ("conditions.space.rh", "공간 상대습도", "제품 설치 공간의 상대습도", "이슬맺힘", "필수", "이슬맺힘", "의뢰자", "숫자", "%", "물리 범위 미정", "예", "Case 공간 RH 열", "현행 통합 정책 없음", "누락 시 차단", "공간 상대습도를 입력해 주세요.", "60%", "빈 값"),
    )
    for item in condition_rows:
        key, label, definition, analysis_types, required, activation, owner, input_type, unit, choices, repeatable, binding, unknown, blocking, question, normal, error = item
        add(
            key, "해석조건", label, definition, analysis_types, required, activation, owner, input_type, unit,
            choices, repeatable, binding, unknown, blocking, question, normal, error, "코드 기준 정리 완료",
            "condition_fieldsets.py; state.py; validator.py; ui.py",
        )
    return rows


def _object_rows() -> tuple[tuple[str, ...], ...]:
    return (
        ("OBJ-CTX", "의뢰 Context", "제품 분류와 해석유형을 확정한 요청 범위", "전체", "1", "1", "아니오", "사업부, 제품군, Platform, 해석유형", "조건 fieldset", "지원하지 않는 유형", "확정 전 미입력", "의뢰자+시스템", "RAC/제품군/Platform/풍량", "코드 기준 정리 완료", "request_context_confirmation.py"),
        ("OBJ-BASE", "Base 제품", "기준 총조립 형상", "전체", "1", "1", "아니오", "geometry_id, 도면번호, 표시명", "형상 축과 Case", "중복 도면번호", "임시 도면번호 허용", "의뢰자+시스템", "형상 1", "코드 기준 정리 완료", "state.py; validator.py"),
        ("OBJ-COMP", "비교 제품", "Base 대비 차이가 있는 추가 총조립 형상", "전체", "0", "제한 없음", "예", "geometry_id, 도면번호, Base 대비 차이", "형상 축과 Case", "Base와 동일 도면번호", "필요 없으면 0개", "의뢰자", "형상 2: 베인 각도 변경", "코드 기준 정리 완료", "state.py; validator.py"),
        ("OBJ-OPER", "운전 조건", "팬 구성과 RPM의 반복 세트", "전체", "1", "제한 없음", "예", "운전명, 팬 개수, 위치, RPM", "Case 운전 열", "현행 물리 금지 조합 없음", "RPM 누락 차단", "의뢰자+시스템", "운전 1", "코드 기준 정리 완료", "condition_fieldsets.py; state.py"),
        ("OBJ-HEX", "열교환기 사양", "열교환기 유형과 형상 파라미터의 반복 세트", "전체", "1", "제한 없음", "예", "사양명, 유형, 관 직경, Fin type, 열 수, FPI", "Case 사양 열", "현행 물리 금지 조합 없음", "활성 필드 누락 차단", "의뢰자+시스템", "사양 1", "코드 기준 정리 완료", "condition_fieldsets.py; heat_exchanger_catalog.py"),
        ("OBJ-SUPPLY", "취출 공기 조건", "취출 온도·습도 반복 세트", "기류 패턴, 이슬맺힘", "1", "제한 없음", "예", "취출 온도, 취출 상대습도", "Case 취출 조건 열", "기류 패턴에서 RH 비활성", "기류 패턴 온도 없음 허용", "의뢰자", "취출 조건 1", "코드 기준 정리 완료", "condition_fieldsets.py"),
        ("OBJ-ROOM", "공간 환경 조건", "설치 공간 온도·습도 반복 세트", "기류 패턴, 이슬맺힘", "1", "제한 없음", "예", "공간 온도, 공간 상대습도", "Case 공간 조건 열", "기류 패턴에서 RH 비활성", "기류 패턴 온도 없음 허용", "의뢰자", "공간 조건 1", "코드 기준 정리 완료", "condition_fieldsets.py"),
        ("OBJ-CASE", "Case", "한 형상과 활성 조건값의 사용자 수동 조합", "전체", "1", "제한 없음", "예", "형상, 운전, 사양, 유형별 온·습도", "미리보기와 Word", "동일 signature 중복 금지", "선택 누락 차단", "의뢰자+시스템", "Case 1", "코드 기준 정리 완료", "state.py; validator.py; case_matrix.py"),
    )


def _case_rule_rows() -> tuple[tuple[str, ...], ...]:
    rows = []
    for analysis_type in EXPECTED_ACTIVE_TYPES:
        columns = get_active_case_matrix_columns({"analysis_type": analysis_type})
        axes = "형상, " + ", ".join(column["label"] for column in columns)
        rows.append((
            f"CASE-{analysis_type}", analysis_type, axes, "각 열에서 하나의 입력 세트를 선택",
            "사용자가 필요한 조합을 별도 행으로 추가", "새 형상당 기본 행 1개",
            "존재하지 않는 형상·조건 선택", "형상과 모든 활성 조건값이 동일한 조합",
            "첫 형상+각 조건의 첫 입력값", "시스템 번호", "해당 없음", "해당 없음",
            f"{analysis_type}: {axes}", "코드 기준 정리 완료", "state.py; validator.py; case_matrix.py",
        ))
    return tuple(rows)


def _term_rows() -> tuple[tuple[str, ...], ...]:
    return (
        ("Context", "사업부·제품군·Platform·해석유형을 묶은 의뢰 범위", "제품 분류, 해석 범위", "조건 Fieldset을 결정하기 위해 필요", "제품 분류와 해석유형을 확정해 주세요.", "모르면 제품 담당자와 확인해 주세요.", "아니오", "예", "확정 전 진행 차단", "제출 차단", "해당 없음", "RAC/제품군/Platform/풍량", "코드 기준 정리 완료", "request_context_confirmation.py"),
        ("Base 제품", "비교의 기준이 되는 첫 번째 총조립 형상", "기준 모델, 기준 형상", "비교 제품의 차이를 설명하기 위해 필요", "기준 제품 도면번호를 입력해 주세요.", "임시 도면번호도 사용할 수 있습니다.", "아니오", "예", "미입력 차단", "제출 차단", "해당 없음", "형상 1", "코드 기준 정리 완료", "ui.py; validator.py"),
        ("운전", "하나 이상의 팬 위치와 RPM을 묶은 조건", "팬 조건, RPM 조건", "Case별 운전 상태를 구분하기 위해 필요", "팬 개수와 RPM을 입력해 주세요.", "여러 팬이면 위치도 구분해 주세요.", "아니오", "예", "RPM 누락 차단", "제출 차단", "해당 없음", "운전 1", "코드 기준 정리 완료", "condition_fieldsets.py; ui.py"),
        ("사양", "열교환기 유형과 상세 값을 묶은 조건", "HEX 사양, 열교환기 조건", "유동 저항과 분포 조건을 구분하기 위해 필요", "열교환기 사양을 입력해 주세요.", "Catalog 값을 확인해 주세요.", "조건부", "예", "필드 누락 차단", "제출 차단", "해당 없음", "사양 1", "코드 기준 정리 완료", "condition_fieldsets.py"),
        ("Case", "형상과 활성 조건을 한 번의 해석 조합으로 연결한 행", "해석 Case, 조건 조합", "실행할 비교 조합을 확정하기 위해 필요", "Case별 형상과 조건을 선택해 주세요.", "입력한 모든 항목을 최소 한 Case에 사용해 주세요.", "아니오", "예", "누락·중복·미사용 항목 검토", "Word 생성 차단", "해당 없음", "Case 1", "코드 기준 정리 완료", "state.py; validator.py"),
        ("없음(온도)", "기류 패턴에서 별도 온도 조건을 적용하지 않는 선택", "온도조건 없음", "온도 미적용 의도를 빈 값과 구분", "별도 온도 조건이 없다면 없음으로 선택해 주세요.", "확인된 온도가 있으면 직접 입력해 주세요.", "아니오", "예", "기류 패턴 온도에만 UI 허용", "안내", "해당 없음", "공간 온도: 없음", "코드 기준 정리 완료", "condition_fieldsets.py; ui.py"),
    )


def _gap_rows() -> tuple[tuple[str, ...], ...]:
    return (
        ("GAP-01", "실제 완료사례", "활성 유형별 최초 요청·재문의·최종 보고서가 저장소에 없음", "높음", "유동해석 담당자", "유형별 정상·다중·정보 부족 사례 3건 수집", "미해결", "02_완료사례_목록"),
        ("GAP-02", "비활성 조건 정의", "material_type, 압력, 설치 환경, 결과 확인 항목 등이 public schema에 있으나 활성 카드에서 사용되지 않음", "중간", "개발자+유동 담당자", "유지·삭제·재활성 여부 결정", "미해결", "constants.py; condition_fieldsets.py"),
        ("GAP-03", "예정 해석유형", "6개 예정 유형은 안내만 있고 Field Rule과 Case 규칙 없음", "중간", "유동해석 담당자", "지원 시 유형별 Master 작성", "미해결", "constants.py"),
        ("GAP-04", "물리 범위", "RPM, 온도, 습도, 열 수, FPI의 물리 허용범위·단위 변환 없음", "높음", "유동해석 담당자", "유형·제품별 범위와 오류 수준 정의", "미해결", "condition_fieldsets.py"),
        ("GAP-05", "모름·해당 없음", "필드 상태는 있으나 조건별 허용과 후속 질문이 통합 Master로 정리되지 않음", "높음", "유동해석 담당자", "06_용어_안내_예외 승인", "미해결", "constants.py; state.py"),
        ("GAP-06", "Case 물리 규칙", "중복과 Coverage 외 물리적 금지 조합·상호 배제 규칙 없음", "중간", "유동해석 담당자", "유형별 금지·필수 조합 정의", "미해결", "state.py; validator.py"),
        ("GAP-07", "팬 위치 검증", "다중 팬 위치는 미리보기에 활용되지만 핵심 조건 validator 필수값이 아님", "중간", "유동해석 담당자", "필수 여부와 오류 수준 결정", "미해결", "condition_fieldsets.py; validator.py; ui.py"),
        ("GAP-08", "결과 판정", "해석유형 안내는 있으나 정량 결과·합격 기준 Master가 없음", "높음", "유동해석 담당자", "유형별 결과 요청과 판정 기준 정의", "미해결", "analysis_type_master.py"),
        ("GAP-09", "초기 사용자값", "특정 부서·요청자·직급이 코드 초기값으로 설정됨", "중간", "서비스 운영자", "로그인 사용자 또는 설정 기반으로 일반화", "미해결", "state.py"),
        ("GAP-10", "유형 Catalog 불일치", "냉매누설은 안내 Master에 있으나 선택 목록에는 없음", "낮음", "유동해석 담당자+개발자", "노출·삭제·향후 상태 결정", "미해결", "analysis_type_master.py; constants.py"),
    )


def _sheet_specs() -> tuple[SheetSpec, ...]:
    common_status = Validation(1, ("코드 기준 정리 완료", "검토 필요", "담당자 승인"))
    return (
        SheetSpec(
            name="작성안내",
            headers=("항목", "내용"),
            rows=(
                ("문서 성격", "현재 h7_v0 코드를 구조해석 SME 회신 기준에 맞춰 역정리한 현행 예시본입니다."),
                ("기준일", "2026-08-20"),
                ("기준 버전", REFERENCE_VERSION),
                ("중요 구분", "코드에 구현된 동작을 정리한 것이며 유동해석 담당자의 업무 재승인을 의미하지 않습니다."),
                ("실제 사례", "저장소에 실제 완료 의뢰·보고서가 없어 코드 동작 예시와 자료 Gap을 구분해 기록했습니다."),
                ("구조 담당자 활용", "현재 유동해석의 분류·필드·객체·Case·검증 수준을 참고하여 구조해석 회신양식에 동일한 수준의 결정을 작성합니다."),
                ("검수 방식", "코드 기준 정리 완료와 담당자 승인을 구분합니다. Gap이 해소되기 전에는 업무 Master 최종 승인으로 간주하지 않습니다."),
                ("코드 근거", "constants.py, condition_fieldsets.py, state.py, validator.py, request_context_confirmation.py, analysis_type_master.py, ui.py"),
            ),
            instruction_sheet=True,
            blank_rows=0,
        ),
        SheetSpec(
            name="01_해석유형_우선순위",
            headers=("사용자 표시명", "주요 질문", "적용 제품/부품", "대표 결과", "적용하지 않는 경우", "수행 빈도", "표준화 가능성", "입력 확보 난이도", "MVP 추천 여부", "현재 UI 상태", "활성 조건 카드", "검수 상태", "코드 근거"),
            rows=_analysis_rows(), reference_sheet=True, blank_rows=0,
        ),
        SheetSpec(
            name="02_완료사례_목록",
            headers=("사례 ID", "해석유형", "자료 성격", "사례 구분", "최초 사용자 요청", "적용 제품/부품", "제품·CAD·도면 정보", "재질 정보", "하중/운전 정보", "구속 정보", "접촉·사양 정보", "담당자 추가 질문", "최종 확정 조건", "최종 Case 목록", "요청 결과", "판정 기준", "최종 의뢰서·보고서 위치", "익명화 상태", "검수 상태", "담당자 의견"),
            rows=_case_reference_rows(), reference_sheet=True, blank_rows=0,
        ),
        SheetSpec(
            name="03_Field_Rule_Master",
            headers=("규칙 ID", "입력 그룹", "사용자 표시명", "정의", "적용 해석유형", "필수 수준", "활성 조건", "입력 주체", "입력 방식", "단위·좌표계", "선택지·범위", "반복 가능 여부", "연결 대상", "모름/해당 없음 처리", "차단 검증", "사용자 질문", "정상 예시", "오류 예시", "검수 상태", "코드 근거"),
            rows=tuple(_common_field_rows()), reference_sheet=True, blank_rows=0,
        ),
        SheetSpec(
            name="04_핵심객체_규칙",
            headers=("객체 규칙 ID", "객체 유형", "업무상 정의", "적용 해석유형", "최소 개수", "최대 개수", "복제 가능 여부", "구성 필드", "연결 대상", "금지 조합", "모름/해당 없음 처리", "입력 주체", "정상 예시", "검수 상태", "코드 근거"),
            rows=_object_rows(), reference_sheet=True, blank_rows=0,
        ),
        SheetSpec(
            name="05_Case_조합규칙",
            headers=("Case 규칙 ID", "적용 해석유형", "Case 구분 축", "동시 조건 여부", "별도 Case 분리 조건", "기본 Case 자동 생성", "금지 조합", "중복 Case 판정 기준", "필수 기준 Case", "Case 이름 입력 주체", "하중 조합계수 입력 주체", "안전계수 입력 주체", "정상 Case 예시", "검수 상태", "코드 근거"),
            rows=_case_rule_rows(), reference_sheet=True, blank_rows=0,
        ),
        SheetSpec(
            name="06_용어_안내_예외",
            headers=("용어/필드", "쉬운 설명", "동의어·사용자 표현", "왜 필요한가", "기본 사용자 질문", "대체 질문", "Agent 추정 허용", "사용자 확인 필요", "모름/해당 없음 처리", "오류 수준", "잘못 선택 시 추천 유형", "도움말·정상 예시", "검수 상태", "코드 근거"),
            rows=_term_rows(), reference_sheet=True, blank_rows=0,
        ),
        SheetSpec(
            name="07_검수_승인",
            headers=("해석유형", "활성 입력 그룹 코드 확인", "필수·조건부 규칙 코드 확인", "객체 연결 코드 확인", "Case 규칙 코드 확인", "모름·예외 업무 승인", "실제 정상 사례 재현", "실제 다중 Case 재현", "실제 입력 누락 사례 재현", "무재문의 착수 판단 가능", "유동해석 담당자", "최종 승인자", "승인일", "검수 상태", "미결 사항·의견"),
            rows=tuple((analysis_type, "예", "예", "예", "예", "아니오", "아니오", "아니오", "아니오", "검토 필요", "", "", "", "검토 필요", "코드 기준 정리는 완료했으나 업무 사례와 SME 승인이 필요") for analysis_type in EXPECTED_ACTIVE_TYPES),
            reference_sheet=True, blank_rows=0,
        ),
        SheetSpec(
            name="08_변경이력",
            headers=("Master 버전", "변경일", "작성자", "변경 이유", "영향 시트", "영향 해석유형", "기존 의뢰서 호환 여부", "검토자", "검수 상태", "비고"),
            rows=((REFERENCE_VERSION, "2026-08-20", "Codex", "h7_v0 현행 코드 기준 최초 역정리", "전체", "전체", "읽기 전용 문서이므로 영향 없음", "", "코드 기준 정리 완료", "애플리케이션 코드/API 변경 없음"),),
            reference_sheet=True, blank_rows=0,
        ),
        SheetSpec(
            name="09_워크숍_기록",
            headers=("회차", "일시", "참석자", "목표", "확정 결정", "미결 사항", "후속 담당자", "완료 예정일", "진행 상태", "비고"),
            rows=(("유동 기준 검수 1차", "미정", "유동해석 담당자, 개발자", "현행 Master와 Gap 검수", "미실시", "실제 사례, 물리 범위, 결과 판정, 모름 처리", "유동해석 담당자", "미정", "예정", "구조해석 워크숍 전에 수행 권장"),),
            reference_sheet=True, blank_rows=0,
        ),
        SheetSpec(
            name="10_현행_Gap_목록",
            headers=("Gap ID", "구분", "현행 상태", "우선도", "결정 담당", "필요 조치", "상태", "코드·시트 근거"),
            rows=_gap_rows(), reference_sheet=True, blank_rows=0,
        ),
    )


def build_xlsx() -> None:
    write_xlsx(XLSX_PATH, _sheet_specs())


def _configure_docx(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    normal = document.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "맑은 고딕"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.font.color.rgb = RGBColor(23, 50, 77)


def build_docx() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document = Document()
    _configure_docx(document)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("h7_v0 현행 유동해석 의뢰 기준서")
    run.bold = True
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(23, 50, 77)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("구조해석 담당자 정보요청 양식의 현행 CFD 작성 예시 · 2026-08-20").italic = True
    document.add_paragraph()

    _add_table(document, ("항목", "내용"), (
        ("기준 버전", REFERENCE_VERSION),
        ("자료 성격", "h7_v0 코드에서 역정리한 현행 동작"),
        ("업무 승인", "미승인 — 유동해석 담당자 검수 필요"),
        ("애플리케이션 변경", "없음"),
    ), (3.6, 12.4))

    document.add_heading("1. 목적과 적용 범위", level=1)
    document.add_paragraph(
        "이 기준서는 구조해석 담당자 회신양식과 같은 관점으로 현재 h7_v0의 유동해석 유형, Field, 반복 객체, "
        "Case와 검증 규칙을 정리한 예시본입니다. 코드에서 확인되지 않는 실제 완료사례, 수행 빈도, 물리적 허용범위와 "
        "합격 기준은 임의로 작성하지 않았습니다."
    )

    document.add_heading("2. 현행 해석유형", level=1)
    active_rows = tuple(
        (analysis_type, RESULT_LABELS[analysis_type], _active_groups(analysis_type))
        for analysis_type in EXPECTED_ACTIVE_TYPES
    )
    _add_table(document, ("선택 가능 유형", "대표 결과", "활성 조건 카드"), active_rows, (4.2, 5.2, 6.6))
    document.add_paragraph("화면상 예정·선택 불가: " + ", ".join(DISABLED_ANALYSIS_TYPE_OPTIONS))
    document.add_paragraph("Legacy 호환용: 일반 유동 해석, 열유동 해석. Catalog 불일치: 냉매누설은 안내 Master에만 존재합니다.")

    document.add_heading("3. 공통 입력 구조", level=1)
    _add_table(document, ("그룹", "핵심 입력", "현행 필수 규칙"), (
        ("의뢰 Context", "사업부, 제품군, Platform, 해석유형", "4개 값이 있어야 Context 확정"),
        ("의뢰자 정보", "사업부, 부서, 요청자, 직급", "모두 필수; 의뢰 번호는 시스템 생성"),
        ("해석 개요", "프로젝트, 개발 등급, NPI, 모델, 날짜, 배경, 확인 내용", "결과 활용 목적을 제외하고 필수"),
        ("제품/형상", "Base 1개, 비교 제품 0개 이상, 도면번호, 차이 설명", "도면번호 형식·중복 검증; 비교 차이 필수"),
        ("해석조건", "운전, 열교환기, 유형별 온·습도", "활성 필드마다 값 필요"),
        ("Case", "형상과 활성 조건의 수동 매핑", "최소 1개; 누락·중복·미사용 항목 검토"),
    ), (3.0, 7.0, 6.0))

    document.add_heading("4. 해석유형별 활성 조건", level=1)
    _add_table(document, ("해석유형", "필수 조건"), (
        ("풍량", "팬 RPM, 열교환기 사양"),
        ("열교환기 유속 프로파일", "팬 RPM, 열교환기 사양"),
        ("기류 패턴", "팬 RPM, 열교환기 사양, 공간 온도, 취출 온도; 온도는 없음 선택 가능"),
        ("이슬맺힘", "팬 RPM, 열교환기 사양, 공간 온도·RH, 취출 온도·RH"),
    ), (4.5, 11.5))
    document.add_paragraph(
        "열교환기 사양은 사양명, 유형, 관 직경(Pi), Fin type, 열 수와 FPI로 구성됩니다. "
        "운전 조건은 한 개 이상의 팬과 각 팬의 RPM으로 구성되며 조건 카드는 반복 추가할 수 있습니다."
    )

    document.add_heading("5. Case Matrix", level=1)
    _add_table(document, ("해석유형", "Case 선택 열"), tuple(
        (
            analysis_type,
            "형상, " + ", ".join(column["label"] for column in get_active_case_matrix_columns({"analysis_type": analysis_type})),
        )
        for analysis_type in EXPECTED_ACTIVE_TYPES
    ), (4.5, 11.5))
    _add_bullets(document, (
        "가능한 모든 조합을 자동 생성하지 않고 사용자가 필요한 조합을 직접 선택합니다.",
        "새 형상에는 각 활성 조건의 첫 번째 값이 연결된 기본 Case가 하나 생성됩니다.",
        "동일한 형상과 활성 조건 조합은 중복으로 판정합니다.",
        "입력한 형상·조건 중 어떤 Case에도 사용되지 않은 항목이 있으면 Coverage 미완료입니다.",
        "Case 누락·중복·Coverage 미완료 상태에서는 Word 의뢰서 생성을 허용하지 않습니다.",
    ))

    document.add_heading("6. 주요 검증", level=1)
    _add_bullets(document, (
        "필수 의뢰자 정보와 해석 개요 누락",
        "Context 미확정, 예정 또는 미지원 해석유형 선택",
        "Base 형상 부재, 도면번호 누락·형식 오류·중복",
        "비교 제품의 Base 대비 차이 누락",
        "활성 조건값 누락",
        "Case 없음, 형상·조건 선택 누락, 동일 Case 중복",
        "Word 생성 시 사용되지 않은 형상·조건 존재",
    ))

    document.add_heading("7. 현행 Gap", level=1)
    _add_table(document, ("구분", "확인이 필요한 내용"), tuple((row[1], row[2]) for row in _gap_rows()), (4.0, 12.0))

    document.add_heading("8. 구조해석 담당자 활용 방법", level=1)
    _add_bullets(document, (
        "Excel의 01~06 시트를 참고하여 구조해석 유형·Field·객체·Case·용어 규칙을 같은 수준으로 작성합니다.",
        "유동 기준의 `코드 기준 정리 완료`는 업무 승인과 다르므로 그대로 복사하지 않습니다.",
        "구조해석의 재질·연결·구속·하중·결과 판정 객체는 현행 CFD에 없으므로 별도로 정의합니다.",
        "유형별 실제 완료사례 3건으로 Master를 재현한 뒤 최종 승인합니다.",
    ))

    document.core_properties.title = "h7_v0 현행 유동해석 의뢰 기준서"
    document.core_properties.subject = "구조해석 담당자 정보요청 양식의 현행 CFD 작성 예시"
    document.core_properties.author = "CAE Request Agent"
    document.save(DOCX_PATH)


def _contract_errors() -> list[str]:
    errors: list[str] = []
    if tuple(ENABLED_ANALYSIS_TYPE_OPTIONS) != EXPECTED_ACTIVE_TYPES:
        errors.append(f"active analysis types changed: {tuple(ENABLED_ANALYSIS_TYPE_OPTIONS)}")
    for analysis_type, expected in EXPECTED_FIELD_KEYS.items():
        actual = _active_fields(analysis_type)
        if actual != expected:
            errors.append(f"active field keys changed for {analysis_type}: {actual}")
    return errors


def check_outputs() -> list[str]:
    errors = _contract_errors()
    if not MARKDOWN_PATH.exists():
        errors.append(f"missing Markdown: {MARKDOWN_PATH}")
    else:
        try:
            text = MARKDOWN_PATH.read_text(encoding="utf-8", errors="strict")
            for phrase in ("현행 사용자 흐름", "Case Matrix 규칙", "현행 Gap"):
                if phrase not in text:
                    errors.append(f"Markdown missing phrase: {phrase}")
        except UnicodeDecodeError as exc:
            errors.append(f"invalid Markdown UTF-8: {exc}")

    try:
        document = Document(DOCX_PATH)
        docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for phrase in ("현행 해석유형", "Case Matrix", "현행 Gap"):
            if phrase not in docx_text:
                errors.append(f"DOCX missing phrase: {phrase}")
    except (OSError, BadZipFile, ValueError) as exc:
        errors.append(f"invalid DOCX: {exc}")

    try:
        with ZipFile(XLSX_PATH) as archive:
            workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
            sheet_names = tuple(sheet.attrib["name"] for sheet in workbook_root.findall(f".//{{{NS_MAIN}}}sheet"))
            if sheet_names != EXPECTED_SHEETS:
                errors.append(f"unexpected XLSX sheets: {sheet_names}")
            for index in range(1, len(EXPECTED_SHEETS) + 1):
                ET.fromstring(archive.read(f"xl/worksheets/sheet{index}.xml"))
    except (OSError, KeyError, BadZipFile, ET.ParseError) as exc:
        errors.append(f"invalid XLSX: {exc}")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description="Build or verify the h7_v0 current CFD reference package.")
    parser.add_argument("--check", action="store_true", help="verify existing outputs without rewriting them")
    args = parser.parse_args()
    if not args.check:
        build_docx()
        build_xlsx()
        print(f"Generated: {DOCX_PATH}")
        print(f"Generated: {XLSX_PATH}")
    errors = check_outputs()
    if errors:
        print("Current CFD reference verification failed:")
        for error in errors:
            print(f"- {error}")
        return 1
    print("Current CFD reference verification passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
